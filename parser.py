import re
from openpyxl import load_workbook


def parse_gs_code(code):
    if not code or not isinstance(code, str):
        raw = str(code).strip() if code else ''
        return '기타', raw, '', raw

    code = code.strip()

    match = re.match(r'^(GS\d{2})(\d{5})([A-Z])(_.+)?$', code)
    if match:
        supplier = match.group(1)
        seq = match.group(2)
        option = match.group(3)
        suffix = match.group(4) or ''
        sku_group = supplier + seq
        return supplier, seq, option + suffix, sku_group

    if code.upper().startswith('GS'):
        match2 = re.match(r'^(GS\d{2})', code)
        if match2:
            supplier = match2.group(1)
            last_alpha = re.search(r'([A-Z])(_.+)?$', code)
            if last_alpha:
                sku_group = code[:last_alpha.start()]
                option = code[last_alpha.start():]
            else:
                sku_group = code
                option = ''
            return supplier, code[4:].rstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ_0123456789') or code[4:], option, sku_group

    return '기타', code, '', code


def parse_excel(file_path):
    wb = load_workbook(file_path, read_only=True, data_only=True)

    result = {
        'products': [],
        'naver_duplicates': [],
        'headers': [],
        'sheet_info': {}
    }

    main_sheet = None
    for name in wb.sheetnames:
        if '전체' in name and '상품' in name:
            main_sheet = wb[name]
            result['sheet_info']['main'] = name
            break
    if not main_sheet:
        main_sheet = wb[wb.sheetnames[0]]
        result['sheet_info']['main'] = wb.sheetnames[0]

    headers = []
    first_row = next(main_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    for i, val in enumerate(first_row):
        headers.append(str(val) if val else f'col_{i+1}')
    result['headers'] = headers

    col_indices = {}
    target_cols = {
        '자체 상품코드': None, '상품코드': None, '상품명': None,
        '판매가': None, '이미지등록(목록)': None, '진열상태': None,
        '판매상태': None, '네이버중복여부': None, '네이버상품번호': None
    }
    for i, h in enumerate(headers):
        if h in target_cols:
            col_indices[h] = i

    for row in main_sheet.iter_rows(min_row=2, values_only=True):
        if not row:
            continue

        idx = col_indices.get('자체 상품코드')
        if idx is None or idx >= len(row) or not row[idx]:
            continue

        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers):
                if val is not None:
                    row_dict[headers[i]] = str(val)
                else:
                    row_dict[headers[i]] = ''

        product_code = str(row[col_indices['자체 상품코드']]).strip()
        supplier, seq, option, sku_group = parse_gs_code(product_code)

        price_str = row_dict.get('판매가', '0')
        try:
            price = float(price_str)
        except (ValueError, TypeError):
            price = 0.0

        result['products'].append({
            'product_code': product_code,
            'cafe24_code': row_dict.get('상품코드', ''),
            'supplier_code': supplier,
            'product_seq': seq,
            'option_code': option,
            'sku_group': sku_group,
            'product_name': row_dict.get('상품명', ''),
            'price': price,
            'image_url': row_dict.get('이미지등록(목록)', ''),
            'display_status': row_dict.get('진열상태', ''),
            'sale_status': row_dict.get('판매상태', ''),
            'naver_status': row_dict.get('네이버중복여부', '신규'),
            'naver_product_id': row_dict.get('네이버상품번호', ''),
            'raw_data': row_dict
        })

    for name in wb.sheetnames:
        if '중복' in name:
            dup_sheet = wb[name]
            dup_headers = []
            first = next(dup_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            for val in first:
                dup_headers.append(str(val) if val else '')

            gs_idx = None
            for i, h in enumerate(dup_headers):
                if 'GS' in h and '코드' in h:
                    gs_idx = i
                    break

            if gs_idx is not None:
                for row in dup_sheet.iter_rows(min_row=2, values_only=True):
                    if row and gs_idx < len(row) and row[gs_idx]:
                        result['naver_duplicates'].append(str(row[gs_idx]).strip())

            result['sheet_info']['duplicates'] = name
            break

    wb.close()
    return result
